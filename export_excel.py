"""Export the researchers shown on the site, with each person's search record, to an Excel workbook.

Reads exactly what the page embeds (build_page.prepare_page with the showcase filter), so the
workbook and the profile cards cannot disagree.
Usage: uvx --with openpyxl python export_excel.py [output.xlsx]
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import build_page

SITE = "https://easonai-5589.github.io/ai-talent-network/"
OUTPUT = build_page.ROOT / "downloads" / "pbcsf-ai-talent-search-evidence.xlsx"
TIERS = {1: "L1 顶级奖项", 2: "L2 高影响力候选", 3: "L3 领域骨干候选", 4: "L4 其他候选"}
POINTS = {1: 100, 2: 85, 3: 60, 4: 35}
HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="264873")
LINK = Font(color="264873", underline="single")
WRAP = Alignment(wrap_text=True, vertical="top")


def table(ws, columns, rows):
    """columns: [(header, width, kind)] with kind in text | int | pct | num | url."""
    ws.append([c[0] for c in columns])
    for cell in ws[1]:
        cell.font, cell.fill, cell.alignment = HEAD, HEAD_FILL, Alignment(wrap_text=True, vertical="center")
    for row in rows:
        ws.append([v[0] if isinstance(v, tuple) else v for v in row])
        r = ws.max_row
        for i, ((_, _, kind), v) in enumerate(zip(columns, row), start=1):
            cell = ws.cell(r, i)
            if kind == "pct" and cell.value is not None:
                cell.number_format = "0.0%"
            elif kind == "num" and cell.value is not None:
                cell.number_format = "0.000"
            elif kind == "url" and v:
                url = v[1] if isinstance(v, tuple) else v
                cell.hyperlink, cell.font = url, LINK
    for i, (_, width, _) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions


def criteria_sheet(ws, merged):
    sr, t = merged["search"], merged["search"]["totals"]
    shown = len(merged["people"])
    seeds = sum(1 for p in merged["people"] if p["search"]["route"] == "人工补录")
    rows = [
        ("检索标准", ""),
        ("数据源", "OpenAlex 开放学术数据库（CC0 许可），覆盖论文、作者和机构"),
        ("检索时间", sr["searched"].replace("/", " 至 ")),
        ("论文范围", f"{sr['since'][:4]} 年 1 月 1 日以来发表的论文"),
        ("学校范围", f"{len(sr['stats'])} 所：亚洲取 QS 2026 亚洲前 20 中的中国大陆、香港、新加坡高校，另加中国科学技术大学；北美取 U.S. News 2026 全美大学前 20（并列第 20 的 4 所全部纳入）；欧洲为牛津和 LSE"),
        ("学校识别", "按 OpenAlex 机构编号检索，包含该校下属学院、研究所和校区"),
        ("检索式 Q1", f"经济金融中使用 AI 的论文：署名机构属于该校，主题领域为“经济学、计量经济学与金融”，且标注“人工智能”或“机器学习”概念。共 {t['q1']:,} 篇"),
        ("检索式 Q2", f"计算机与经济金融交叉论文：署名机构属于该校，主题领域同时包含“计算机科学”与“经济学、计量经济学与金融”。共 {t['q2']:,} 篇"),
        ("作者归属", "只计该论文中署名机构属于该校的作者；同一作者在该校的 Q1、Q2 论文合并去重，称为“交叉论文”"),
        ("AI 论文判定", "标题含机器学习、深度学习、大语言模型、文本分析、预测等方法词，或论文主题属于人工智能、计算机视觉等子领域"),
        ("金融等级判定", "标题和主题含金融词（资产、银行、信贷、交易、加密货币、会计等）计 1；含一般经济词或属经济金融领域计 0.7；其余计 0.35"),
        ("现职检索", f"逐人检索学校官方页面（院系教师名录、个人主页），核对姓名、学校、职称和院系，记录来源网址和核验日期；最近一轮完成于 {merged.get('role_checked_at')}"),
        ("人工补录", f"OpenAlex 单位记错而漏检的学者按论文署名补录，共 {seeds} 位"),
        ("", ""),
        ("入库门槛（6 项同时满足）", ""),
        ("交叉论文", "以该校署名发表至少 2 篇交叉论文"),
        ("当前关联", "OpenAlex 最近关联机构包含该校"),
        ("研究资历", "h 指数 ≥ 6"),
        ("领域对口", "医学、环境、物理等非目标领域的主题占比 < 40%"),
        ("经济金融背景", "经济金融相关主题占比 ≥ 15%；计算机背景的学者还须有 ≥ 3 篇交叉论文"),
        ("AI 背景", "AI 论文 ≥ 2 篇，或 AI 主题占比 ≥ 15%，或计算机主题占比 ≥ 20% 且交叉论文 ≥ 3 篇"),
        ("", ""),
        ("评价维度", ""),
        ("学术影响力（L1–L4）", "L1 诺贝尔经济学奖或图灵奖；L2 克拉克奖、奈望林纳奖等，或经济金融 h≥45/被引≥1.5 万、计算机与交叉 h≥60/被引≥2.5 万，或人工标注领域带头人；L3 经济金融 20≤h<45、计算机与交叉 25≤h<60；L4 其余。分层分 L1 100、L2 85、L3 60、L4 35"),
        ("AI × 金融相关性", "100 × (0.55 × 交叉论文得分 + 0.25 × 金融主题得分 + 0.20 × AI 强度得分)。交叉论文得分 = 1 − e^(−E/5)，E 为每篇交叉论文金融等级之和（非 AI 论文乘 0.4）；金融主题占比 35% 即满分；AI 强度取 AI 主题占比与 0.3 × AI 论文数/6 的较大值，30% 即满分"),
        ("综合分", "0.6 × 相关性 + 0.4 × 分层分"),
        ("网站收录", "只收录现职经学校官方页面核实的学者"),
        ("", ""),
        ("检索漏斗（按学校计次）", ""),
        ("Q1 + Q2 去重后论文", f"{t['works']:,} 篇"),
        ("本校署名作者", f"{t['authors']:,} 人次"),
        ("交叉论文 ≥ 2 篇的候选作者", f"{t['candidates']:,} 人次"),
        ("未通过（当前关联 / h / 领域 / 经济金融 / AI）", f"{t['not_current']:,} / {t['h_lt6']:,} / {t['off_target']:,} / {t['no_econ']:,} / {t['no_ai']:,}"),
        ("通过全部门槛", f"{t['kept']:,} 人次"),
        ("合并重复档案、加人工补录后的候选", f"{merged['showcase']['pool']:,} 位"),
        ("现职经官方核实，网站收录", f"{shown:,} 位"),
        ("", ""),
        ("网站", SITE),
    ]
    for k, v in rows:
        ws.append([k, v])
        if k and not v:
            ws.cell(ws.max_row, 1).font = Font(bold=True, size=12, color="264873")
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 120
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = WRAP
    ws.cell(ws.max_row, 2).hyperlink, ws.cell(ws.max_row, 2).font = SITE, LINK


def people_rows(merged):
    people = sorted(merged["people"], key=lambda p: (-p["score"], -p["fit"], p["name"]))
    stats = {s["school"]: s["school_zh"] for s in merged["search"]["stats"]}
    for i, p in enumerate(people, start=1):
        s, role = p["search"], p.get("display_role") or p.get("current_role") or {}
        sh = s["shares"]
        yield p, [
            i, p["name"], p.get("name_zh") or "", role.get("institution") or "", role.get("position") or "", role.get("department") or "", p.get("region") or "",
            s["route"] + (f"：{s['route_why']}" if s.get("route_why") else ""), stats.get(s["school"], s["school"]), "；".join(s.get("matched_institutions") or []) or "人工补录，按论文署名确认",
            s["q1_papers"], s["q2_papers"], s["cross"], s["ai_papers"], p["h"], p["cites"],
            sh.get("off"), sh.get("econ"), sh.get("cs"), sh.get("ai"), "；".join(s["ai_gate"]) or "人工补录",
            "经济金融" if s["home"] == "金融/经济" else "计算机 / 交叉", TIERS.get(p["tier"], ""), s["tier_why"], POINTS.get(p["tier"]),
            s["E"], s["cross_score"], s["fin_share"], s["fin_score"], s["ai_eff"], s["ai_score"], p["fit"], p["score"], "是" if p.get("shortlist") else "",
            "已官方核实", role.get("source_url") or "", role.get("checked_at") or "",
            f"https://openalex.org/{p['id']}", p.get("orcid") or "", "; ".join(p.get("profile_urls") or []),
            ("打开档案卡", f"{SITE}#p-{p['id']}"),
        ]


PEOPLE_COLUMNS = [
    ("序号", 6, "int"), ("姓名", 22, "text"), ("中文名", 9, "text"), ("现任学校 / 机构", 20, "text"), ("职称", 16, "text"), ("院系", 26, "text"), ("地区", 6, "text"),
    ("检索途径", 12, "text"), ("检索学校", 14, "text"), ("署名机构（当前关联）", 28, "text"),
    ("Q1 论文", 8, "int"), ("Q2 论文", 8, "int"), ("交叉论文", 8, "int"), ("AI 论文", 8, "int"), ("h 指数", 7, "int"), ("总被引", 9, "int"),
    ("非目标领域占比（<40%）", 11, "pct"), ("经济金融主题占比（≥15%）", 11, "pct"), ("计算机主题占比", 10, "pct"), ("AI 主题占比", 10, "pct"), ("AI 背景（满足条件）", 30, "text"),
    ("学科口径", 11, "text"), ("分层", 15, "text"), ("分层依据", 34, "text"), ("分层分", 7, "int"),
    ("交叉论文证据 E", 9, "num"), ("交叉论文得分", 9, "pct"), ("金融主题占比", 9, "pct"), ("金融主题得分", 9, "pct"), ("AI 强度", 9, "pct"), ("AI 强度得分", 9, "pct"),
    ("AI × 金融相关性", 9, "int"), ("综合分", 7, "int"), ("Shortlist", 8, "text"),
    ("现职核验", 10, "text"), ("现职来源", 36, "url"), ("核验日期", 11, "text"),
    ("OpenAlex 档案", 32, "url"), ("ORCID", 32, "url"), ("个人主页", 36, "text"), ("网站档案卡", 12, "url"),
]

PAPER_COLUMNS = [
    ("序号", 6, "int"), ("学者", 22, "text"), ("中文名", 9, "text"), ("检索学校", 14, "text"), ("论文标题", 60, "text"), ("年份", 6, "int"), ("被引", 7, "int"),
    ("命中检索式", 10, "text"), ("AI 论文", 8, "text"), ("AI 判定依据", 28, "text"), ("金融等级", 8, "num"), ("金融等级依据", 26, "text"), ("链接", 40, "url"),
]


def export(output=OUTPUT):
    _, merged = build_page.prepare_page(showcase=True)
    if "search" not in merged:
        raise SystemExit("data/search_evidence.json is required for the export")
    wb = Workbook()
    criteria_sheet(wb.active, merged)
    wb.active.title = "检索标准"
    rows = list(people_rows(merged))
    table(wb.create_sheet(f"学者（{len(rows)}）"), PEOPLE_COLUMNS, [r for _, r in rows])
    papers, n = [], 0
    stats = {s["school"]: s["school_zh"] for s in merged["search"]["stats"]}
    for p, _ in rows:
        s = p["search"]
        for oa, ai, q in s["papers"]:
            w = merged["papers"][oa]
            n += 1
            papers.append([n, p["name"], p.get("name_zh") or "", stats.get(s["school"], s["school"]), w["title"], w["year"], w["cites"],
                           q or "补录", "是" if ai else "否（×0.4）", w["ai_why"] or "", w["fin"], w["fin_why"], w["doi"] or f"https://openalex.org/{oa}"])
    table(wb.create_sheet("参考论文"), PAPER_COLUMNS, papers)
    Path(output).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return {"people": len(rows), "papers": len(papers), "output": str(output)}


if __name__ == "__main__":
    print(export(sys.argv[1] if len(sys.argv) > 1 else OUTPUT))
